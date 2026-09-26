"""
NORDIC NEXUS alerts for the Tanzania Tender Scanner (Nordic Construction Company Ltd)
Developed by Eng. Frank and Eng. Michael.

After every hourly scan this sends ONE digest per person per channel containing:
  - NEW tenders (Pursue / Check) found since the last digest
  - DEADLINE reminders 7, 3 and 1 day(s) before submission for open Pursue / Check tenders
by email (company mail server or Outlook), SMS (Beem Africa) and WhatsApp (WhatsApp Business
Cloud API). Each item is sent once only (alerts_sent.json).

Who receives what:
  recipients.json – written automatically from the NEXUS Staff page by the NEXUS sync
                    (email, phone and each person's alert choices)
  scanner_config.json → alerts → recipients – extra people not in NEXUS (optional)

It also creates a folder per new Pursue / Check tender under "Tender Documents", with a
shortcut to the NeST page, so staff save the downloaded tender document in one place.

SECURITY: no password, API key or token is written in this folder. They are kept in
Windows Credential Manager (encrypted, tied to this Windows account) and saved with
"3 - Save alert keys.bat".
"""
import base64, json, os, re, smtplib, ssl, urllib.request, urllib.error
from datetime import datetime
from email.message import EmailMessage

HERE = os.path.dirname(os.path.abspath(__file__))
SVC_SMTP = "NORDIC-NEXUS-SMTP"
SVC_BEEM = "NORDIC-NEXUS-Beem"
SVC_WA = "NORDIC-NEXUS-WhatsApp"
OPEN = ("Watching", "Preparing")
NEXUS_URL = "https://claude.ai/artifact/AZpbRc76i8iSHEr3paoW92"


def _log(msg):
    print(f"{datetime.now():%Y-%m-%d %H:%M:%S}  {msg}", flush=True)


def _load(name, default):
    try:
        with open(os.path.join(HERE, name), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return default


def _save(name, data):
    try:
        with open(os.path.join(HERE, name), "w", encoding="utf-8") as f:
            json.dump(data, f, indent=1, ensure_ascii=False)
    except OSError as e:
        _log(f"Could not save {name}: {e}")


ENV_KEYS = {(SVC_SMTP, None): "SMTP_PASSWORD", (SVC_BEEM, "api_key"): "BEEM_API_KEY",
            (SVC_BEEM, "secret_key"): "BEEM_SECRET_KEY", (SVC_WA, "token"): "WHATSAPP_TOKEN"}


def secret(service, key):
    """Cloud (GitHub Actions): keys come from GitHub Secrets as environment variables.
    Office PC: keys come from Windows Credential Manager."""
    env = ENV_KEYS.get((service, key)) or ENV_KEYS.get((service, None))
    if env and os.environ.get(env):
        return os.environ[env]
    try:
        import keyring
        return keyring.get_password(service, key)
    except Exception:
        return None


def safe_name(s, n=70):
    return re.sub(r"[^A-Za-z0-9 ._()\-]+", "-", str(s or "")).strip(" .-")[:n] or "tender"


def _abs(path):
    path = str(path).replace("\\", "/").replace("/", os.sep)
    return path if os.path.isabs(path) else os.path.normpath(os.path.join(HERE, path))


def tz_phone(p):
    """+255 7xx … in any common local format → 2557xxxxxxxx (digits only, as the APIs want)."""
    d = re.sub(r"\D", "", str(p or ""))
    if d.startswith("0") and len(d) == 10:
        d = "255" + d[1:]
    if len(d) == 9 and d[0] in "67":
        d = "255" + d
    return d if len(d) >= 11 else ""


# ------------------------------------------------------------------ recipients
def recipients(acfg):
    """NEXUS Staff (recipients.json, refreshed by the sync) + extra people in the config."""
    people = []
    env = os.environ.get("RECIPIENTS_JSON")
    try:
        src = json.loads(env) if env else _load("recipients.json", {})
    except ValueError:
        src = {}
    staff = src if isinstance(src, list) else src.get("staff", [])
    for s in staff:
        kinds = {"All tender alerts": ["new", "deadline"], "Deadlines only": ["deadline"], "None": []}.get(
            s.get("alert_types") or "All tender alerts", ["new", "deadline"])
        if not kinds or s.get("status") == "Left":
            continue
        people.append({"name": s.get("name", ""), "email": s.get("email", ""), "phone": s.get("phone", ""),
                       "by_email": s.get("n_email", "Yes") == "Yes", "by_sms": s.get("n_sms") == "Yes",
                       "by_whatsapp": s.get("n_whatsapp") == "Yes", "gets": kinds, "active": True})
    seen = {(p["email"] or "").lower() for p in people if p["email"]} | {tz_phone(p["phone"]) for p in people if p["phone"]}
    for p in acfg.get("recipients", []):
        if (p.get("email") or "").lower() in seen or (p.get("phone") and tz_phone(p["phone"]) in seen):
            continue
        people.append(p)
    return [p for p in people if p.get("active", True)]


# ------------------------------------------------------------------ read open tenders
def open_tenders(workbook_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = wb["Pipeline"].iter_rows(min_row=4, values_only=True)
    hdr = [str(h or "").strip() for h in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in rows:
        def g(h):
            return r[ix[h]] if h in ix and ix[h] < len(r) else None
        if not g("Tender no.") and not g("Tender title"):
            continue
        out.append({"no": str(g("Tender no.") or ""), "title": str(g("Tender title") or ""), "client": str(g("Client") or ""),
                    "location": str(g("Location") or ""), "deadline": g("Submission date"), "status": str(g("Status") or ""),
                    "screen": str(g("Initial screen") or ""), "url": str(g("Source URL") or ""),
                    "key": str(g("Source ID") or g("Tender no.") or "")})
    wb.close()
    return [t for t in out if t["status"] in OPEN and t["screen"] in ("Pursue", "Check")]


# ------------------------------------------------------------------ messages
def _when(d):
    return d.strftime("%a %d %b %H:%M") if isinstance(d, datetime) else "see notice"


def compose(new, due):
    """new: [tender]; due: [(tender, days_left)] → (subject, long text, html, short SMS text)"""
    parts = []
    if new:
        parts.append(f"{len(new)} new tender{'s' if len(new) != 1 else ''}")
    if due:
        parts.append(f"{len(due)} deadline{'s' if len(due) != 1 else ''} coming")
    subject = "NORDIC NEXUS: " + ", ".join(parts)
    L = [subject, ""]
    if due:
        L.append("CLOSING SOON")
        L += [f"- {d} day{'s' if d != 1 else ''} left: {t['title'][:95]} | {t['client'][:40]} | closes {_when(t['deadline'])} | {t['no']}" for t, d in due]
        L.append("")
    if new:
        L.append("NEW TENDERS")
        L += [f"- [{t['screen']}] {t['title'][:95]} | {t['client'][:40]} | {t.get('location', '')[:30]} | closes {_when(t['deadline'])}" for t in new]
        L.append("")
    L.append(f"Open NORDIC NEXUS → Tender Radar for details and AI bid advice: {NEXUS_URL}")
    text = "\n".join(L)
    row = lambda a, b, c, d: f"<tr><td style='padding:6px 8px;border-bottom:1px solid #D4DAE3'>{a}</td><td style='padding:6px 8px;border-bottom:1px solid #D4DAE3'><b>{b}</b><br><span style='color:#5A6580'>{c}</span></td><td style='padding:6px 8px;border-bottom:1px solid #D4DAE3;white-space:nowrap'>{d}</td></tr>"
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html = (f"<div style='font-family:Segoe UI,Arial,sans-serif;color:#182340;max-width:640px'>"
            f"<div style='background:#182340;color:#fff;padding:14px 18px;border-top:6px solid #F2A900'><b style='font-size:18px;letter-spacing:1px'>NORDIC NEXUS</b><br>"
            f"<span style='color:#9FC7EC'>{esc(subject[14:])}</span></div>")
    if due:
        html += "<h3 style='margin:16px 0 6px'>Closing soon</h3><table style='border-collapse:collapse;width:100%;font-size:13px'>" + "".join(
            row(f"<b style='color:{'#B8402A' if d <= 3 else '#9A6500'}'>{d}d</b>", esc(t['title']), esc(t['client']) + ' · ' + esc(t['no']), esc(_when(t['deadline']))) for t, d in due) + "</table>"
    if new:
        html += "<h3 style='margin:16px 0 6px'>New tenders</h3><table style='border-collapse:collapse;width:100%;font-size:13px'>" + "".join(
            row(esc(t['screen']), esc(t['title']), esc(t['client']) + (' · ' + esc(t.get('location', '')) if t.get('location') else ''), esc(_when(t['deadline']))) for t in new) + "</table>"
    html += (f"<p style='margin-top:18px'><a href='{NEXUS_URL}' style='background:#539FDC;color:#0E1830;padding:9px 14px;text-decoration:none;font-weight:bold'>Open Tender Radar</a></p>"
             f"<p style='font-size:11px;color:#5A6580'>Nordic Construction Company Ltd · automatic alert from NORDIC NEXUS · change what you receive on the Staff page</p></div>")
    first = (sorted(due, key=lambda x: x[1])[0] if due else None)
    sms = "NORDIC NEXUS: " + ", ".join(parts) + "."
    if first:
        sms += f" Next: {first[0]['title'][:60]} closes {_when(first[0]['deadline'])}."
    elif new:
        sms += f" e.g. {new[0]['title'][:70]}."
    sms += " See Tender Radar."
    return subject, text, html, sms[:320]


# ------------------------------------------------------------------ channels
def send_email(acfg, to, subject, text, html):
    ecfg = acfg.get("email", {})
    if ecfg.get("method", "smtp") == "outlook":
        import win32com.client  # needs Outlook signed in on this PC
        m = win32com.client.Dispatch("Outlook.Application").CreateItem(0)
        m.To, m.Subject, m.HTMLBody = to, subject, html
        m.Send()
        return
    user = ecfg.get("smtp_user", "")
    pw = secret(SVC_SMTP, user)
    if not user or not pw:
        raise RuntimeError("mailbox or password not saved – run '3 - Save alert keys.bat', option 1")
    msg = EmailMessage()
    msg["From"] = f"NORDIC NEXUS <{ecfg.get('from') or user}>"
    msg["To"], msg["Subject"] = to, subject
    msg.set_content(text)
    msg.add_alternative(html, subtype="html")
    host, port = ecfg.get("smtp_host", "mail.nordictz.co.tz"), int(ecfg.get("smtp_port", 465))
    ctx = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctx, timeout=30) as s:
            s.login(user, pw); s.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=30) as s:
            s.starttls(context=ctx); s.login(user, pw); s.send_message(msg)


def _post(url, payload, headers):
    req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), method="POST",
                                 headers={"Content-Type": "application/json", **headers})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.read().decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"HTTP {e.code}: {e.read().decode('utf-8', 'replace')[:200]}")


def send_sms(acfg, phones, text):
    scfg = acfg.get("sms", {})
    key, sec = secret(SVC_BEEM, "api_key"), secret(SVC_BEEM, "secret_key")
    if not key or not sec:
        raise RuntimeError("Beem API key not saved – run '3 - Save alert keys.bat', option 2")
    auth = base64.b64encode(f"{key}:{sec}".encode()).decode()
    body = {"source_addr": scfg.get("sender_id", "INFO"), "schedule_time": "", "encoding": 0, "message": text,
            "recipients": [{"recipient_id": i + 1, "dest_addr": p} for i, p in enumerate(phones)]}
    _post(scfg.get("url", "https://apisms.beem.africa/v1/send"), body, {"Authorization": "Basic " + auth})


def send_whatsapp(acfg, phone, text):
    wcfg = acfg.get("whatsapp", {})
    token = secret(SVC_WA, "token")
    if not token or not wcfg.get("phone_number_id"):
        raise RuntimeError("WhatsApp token or phone_number_id missing – see the Setup Guide")
    url = f"https://graph.facebook.com/{wcfg.get('api_version', 'v21.0')}/{wcfg['phone_number_id']}/messages"
    body = {"messaging_product": "whatsapp", "to": phone, "type": "template",
            "template": {"name": wcfg.get("template", "tender_alert"), "language": {"code": wcfg.get("language", "en")},
                         "components": [{"type": "body", "parameters": [{"type": "text", "text": text.replace("\n", " · ")[:1000]}]}]}}
    _post(url, body, {"Authorization": "Bearer " + token})


def deliver(acfg, kind_new, kind_due, subject, text, html, sms, log=_log):
    """One digest per person per channel. Returns {'email': n, 'sms': n, 'whatsapp': n}."""
    sent = {"email": 0, "sms": 0, "whatsapp": 0}
    want = lambda p: (kind_new and "new" in p.get("gets", [])) or (kind_due and "deadline" in p.get("gets", []))
    people = [p for p in recipients(acfg) if want(p)]
    sms_to = []
    for p in people:
        if p.get("email") and p.get("by_email", True) and acfg.get("email", {}).get("enabled"):
            try:
                send_email(acfg, p["email"], subject, text, html); sent["email"] += 1
            except Exception as e:
                log(f"Email to {p.get('name')} failed: {e}")
        ph = tz_phone(p.get("phone"))
        if ph and p.get("by_whatsapp") and acfg.get("whatsapp", {}).get("enabled"):
            try:
                send_whatsapp(acfg, ph, text); sent["whatsapp"] += 1
            except Exception as e:
                log(f"WhatsApp to {p.get('name')} failed: {e}")
        if ph and p.get("by_sms") and acfg.get("sms", {}).get("enabled"):
            sms_to.append(ph)
    if sms_to:
        try:
            send_sms(acfg, sorted(set(sms_to)), sms); sent["sms"] = len(set(sms_to))
        except Exception as e:
            log(f"SMS failed: {e}")
    return sent


# ------------------------------------------------------------------ tender folders
def make_folders(cfg, tenders, log):
    root = cfg.get("documents_folder")
    if not root:
        return
    root = _abs(root)
    for t in tenders:
        d = os.path.join(root, safe_name(f"{t['client']} - {t['no']}"))
        if os.path.isdir(d):
            continue
        try:
            os.makedirs(d, exist_ok=True)
            with open(os.path.join(d, "Open on NeST.url"), "w", encoding="utf-8") as f:
                f.write(f"[InternetShortcut]\nURL={t['url'] or 'https://nest.go.tz/tenders/published-tenders'}\n")
            with open(os.path.join(d, "README.txt"), "w", encoding="utf-8") as f:
                f.write(f"{t['title']}\nClient: {t['client']}\nTender no.: {t['no']}\nCloses: {_when(t['deadline'])}\n\n"
                        "Log in to NeST with Nordic's tenderer account, download the tender document and save it here.\n"
                        "Then attach it to this tender in NORDIC NEXUS so the whole team uses one copy.\n")
        except OSError as e:
            log(f"Could not create folder {d}: {e}")


# ------------------------------------------------------------------ entry point
def write_feed(workbook_path, log=_log):
    """data/nexus_feed.json – PUBLIC tender data ready for the NEXUS sync task (no code needed there)."""
    import nexus_sync as NS
    rows, out, seen = NS.read_workbook(workbook_path), [], set()
    for o in rows:
        i = NS.tender_id(o)
        if i in seen:
            continue
        seen.add(i)
        out.append({"id": i, "title": o.get("title", ""), "deadline": o.get("deadline", ""),
                    "body": {**o, "status": o.get("status") or "Watching", "synced_by": "NEXUS auto-sync"}})
    feed = {"generated": datetime.now().strftime("%Y-%m-%dT%H:%M:%S") + "+03:00",
            "scanner": NS.scanner_status(os.path.join(HERE, "last_scan.txt")),
            "workbook_rows": len(rows), "tenders": out}
    with open(_abs("../data/nexus_feed.json"), "w", encoding="utf-8") as f:
        json.dump(feed, f, ensure_ascii=False, indent=0)


def after_scan(cfg, workbook_path, fresh, now=None, log=_log):
    try:
        write_feed(workbook_path, log)
    except Exception as e:
        log(f"NEXUS feed not written: {e}")
    acfg = cfg.get("alerts", {})
    now = now or datetime.now()
    try:
        tenders = open_tenders(workbook_path)
    except Exception as e:
        log(f"Alerts: could not read the workbook ({e})"); tenders = []
    new = [{"key": str(t.get("key")), "no": t.get("no") or "", "title": t.get("title", ""), "client": t.get("client", ""),
            "location": t.get("loc", ""), "deadline": t.get("deadline"), "screen": t.get("screen"), "url": t.get("source_url", "")}
           for t in fresh if t.get("screen") in ("Pursue", "Check")]
    make_folders(cfg, new + [t for t in tenders if t["screen"] == "Pursue"], log)
    if not acfg.get("enabled"):
        return
    sent = _load("alerts_sent.json", {})
    new = [t for t in new if "new|" + t["key"] not in sent]
    due, due_keys = [], []
    days_list = sorted(acfg.get("deadline_days", [7, 3, 1]))
    for t in tenders:
        if not isinstance(t["deadline"], datetime):
            continue
        left = (t["deadline"] - now).total_seconds() / 86400
        win = next((d for d in days_list if 0 < left <= d), None)
        if win is None:
            continue
        k = f"deadline|{t['key']}|{win}"
        if k in sent:
            continue
        due.append((t, max(1, int(left + 0.999)))); due_keys.append(k)
        for bigger in days_list:
            if bigger > win:
                sent.setdefault(f"deadline|{t['key']}|{bigger}", "skipped")
    if not new and not due:
        _save("alerts_sent.json", sent); return
    subject, text, html, sms = compose(new, sorted(due, key=lambda x: x[1]))
    n = deliver(acfg, bool(new), bool(due), subject, text, html, sms, log)
    stamp = f"{now:%Y-%m-%d %H:%M}"
    for t in new:
        sent["new|" + t["key"]] = stamp
    for k in due_keys:
        sent[k] = stamp
    _save("alerts_sent.json", sent)
    log(f"Alerts sent – {subject[14:]}: {n['email']} email, {n['sms']} SMS, {n['whatsapp']} WhatsApp")
