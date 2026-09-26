"""
NORDIC NEXUS sync – turns the scanner's Pipeline workbook into database writes for NEXUS.
Developed by Eng. Frank and Eng. Michael.

Used by the automatic NEXUS sync (a scheduled Claude task). It reads, never changes, the workbook:

  python nexus_sync.py --workbook Nordic_Opportunity_Pipeline.xlsx --db-dir <folder with tenders/*.json>
                       [--last-scan last_scan.txt] --out <folder>

<db-dir>/tenders/<id>.json are the current NEXUS tender documents (as exported by ArtifactData
list with out_dir). The script writes <out>/writes_NN.json (batches of at most 50 writes, each
{"op","collection","doc_id","file_path"}) plus <out>/docs/*.json bodies, and prints a summary.

Rules (same as the manual "Import pipeline" button):
  - a tender not yet in NEXUS is added in full
  - for a tender already in NEXUS only EMPTY fields are filled, and the submission deadline is
    updated when the source changed it; edits made in NEXUS (status, notes, scores, AI advice)
    are never overwritten
"""
import argparse, glob, json, os, re
# v2 (26 Sep 2026): stable ids for notices without a tender number; tender updates carry if_version from versions.json
from datetime import datetime, date, timedelta, timezone

XMAP = {'Tender no.': 'no', 'Date found': 'date_found', 'Source': 'source', 'Source URL': 'url', 'Client': 'client',
        'Client type': 'client_type', 'Tender title': 'title', 'Scope summary': 'scope', 'Location': 'location',
        'Category': 'category', 'CRB class required': 'crb', 'Est. value TZS': 'value', 'Funding confirmed': 'funding',
        'Site visit date': 'site_visit', 'Clarification deadline': 'clar', 'Submission date': 'deadline',
        'Bid security': 'bid_security', 'Documents purchased': 'docs', 'Status': 'status', 'Owner': 'owner',
        'Outcome': 'outcome', 'Lessons': 'lessons', 'Date advertised': 'advertised', 'Notice type': 'notice',
        'Initial screen': 'screen', 'Screen reason': 'reason', 'Next action': 'next', 'Found by': 'found_by',
        'Source ID': 'source_id', 'MD decision': 'md_decision'}
TZ = "+03:00"   # the workbook holds East Africa Time without a zone


def safe_id(s):
    return re.sub(r"[^A-Za-z0-9_\-.~:@+]", "-", str(s))[:120]


def iso(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%S") + TZ
    if isinstance(v, date):
        return v.strftime("%Y-%m-%dT00:00:00") + TZ
    return v


def same_time(a, b):
    try:
        pa = datetime.fromisoformat(str(a).replace("Z", "+00:00"))
        pb = datetime.fromisoformat(str(b).replace("Z", "+00:00"))
        return abs((pa - pb).total_seconds()) < 60
    except ValueError:
        return str(a) == str(b)


def read_workbook(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Pipeline"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and "Tender no." in r)
    hdr = rows[hi]
    out = []
    for r in rows[hi + 1:]:
        o = {}
        for h, v in zip(hdr, r):
            k = XMAP.get(h)
            if not k or v is None or v == "" or (isinstance(v, str) and v.startswith("=")):
                continue
            if k == "value" and not isinstance(v, (int, float)):
                continue
            o[k] = iso(v)
        if o.get("no") or o.get("title"):
            out.append(o)
    wb.close()
    return out


PLACEHOLDER_NO = re.compile(r"^\s*(\(.*\)|n/?a|tbc|tba|none|not shown|-+|see .*)\s*$", re.I)


def tender_id(o):
    """Stable NEXUS id: the tender number, or – when the notice shows none – the scanner's source id."""
    no = str(o.get("no") or "").strip()
    if no and not PLACEHOLDER_NO.match(no):
        return safe_id(no)
    if o.get("source_id"):
        return safe_id(o["source_id"])
    return safe_id((o.get("client", "") + " " + o.get("title", "")).strip() or no)


def read_versions(path):
    """versions.json: {"<doc_id>": <version>, ...} as shown by the ArtifactData listing."""
    if path and os.path.exists(path):
        try:
            return {str(k): int(v) for k, v in json.load(open(path, encoding="utf-8")).items() if v}
        except (ValueError, TypeError, AttributeError):
            pass
    return {}


def read_db(db_dir, versions=None):
    versions = versions or {}
    docs = {}
    for f in glob.glob(os.path.join(db_dir, "tenders", "*.json")):
        with open(f, encoding="utf-8") as fh:
            d = json.load(fh)
        body = d.get("data", d)            # accept {"id","data","version"} or a bare body
        doc_id = d.get("id") or os.path.splitext(os.path.basename(f))[0]
        docs[doc_id] = {"data": body, "version": d.get("version") or versions.get(doc_id)}
    return docs


def scanner_status(path):
    if not path or not os.path.exists(path):
        return None
    txt = open(path, encoding="utf-8", errors="replace").read()
    first = txt.splitlines()[0] if txt else ""
    m = re.match(r"(\d{4}-\d\d-\d\d \d\d:\d\d)", first)
    sources = []
    for line in txt.splitlines():
        s = re.match(r"\s+(.+?): last OK ([^;]+)(?:; last error (.+))?$", line)
        if s:
            sources.append({"name": s.group(1), "last_ok": s.group(2).strip(), "error": (s.group(3) or "").strip()[:200]})
    return {"last_scan": (m.group(1).replace(" ", "T") + ":00" + TZ) if m else None, "summary": first[18:400], "sources": sources}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--db-dir", required=True)
    ap.add_argument("--last-scan")
    ap.add_argument("--out", required=True)
    ap.add_argument("--versions", help="JSON file {doc_id: version} for the exported tenders (default <db-dir>/versions.json)")
    ap.add_argument("--status-version", type=int, help="current version of system/sync in NEXUS (needed to overwrite it)")
    ap.add_argument("--staff-dir", help="folder with staff/*.json exported from NEXUS – writes <out>/recipients.json for the alerts")
    a = ap.parse_args()
    now = (datetime.now(timezone.utc) + timedelta(hours=3)).strftime("%Y-%m-%dT%H:%M:%S") + TZ   # EAT, whatever the machine clock zone
    os.makedirs(os.path.join(a.out, "docs"), exist_ok=True)
    for f in glob.glob(os.path.join(a.out, "writes_*.json")) + glob.glob(os.path.join(a.out, "docs", "*.json")):
        os.remove(f)

    versions = read_versions(a.versions or os.path.join(a.db_dir, "versions.json"))
    wb_rows, db = read_workbook(a.workbook), read_db(a.db_dir, versions)
    missing_versions = []
    writes, added, updated, deadline_changes = [], [], [], []
    seen = set()
    for o in wb_rows:
        doc_id = tender_id(o)
        if doc_id in seen:
            continue
        seen.add(doc_id)
        ex = db.get(doc_id)
        if ex is None:
            body = {**o, "status": o.get("status") or "Watching", "imported_at": now, "synced_by": "NEXUS auto-sync"}
            op = {"op": "set"}
            added.append(o.get("title", doc_id))
        else:
            cur = ex["data"]
            body = {}
            for k, v in o.items():
                if k == "deadline":
                    if cur.get("deadline") and not same_time(cur["deadline"], v):
                        body[k] = v
                        body["deadline_note"] = f"Deadline changed at source (was {cur['deadline'][:16].replace('T', ' ')})"
                        deadline_changes.append(o.get("no") or doc_id)
                    elif not cur.get("deadline"):
                        body[k] = v
                elif cur.get(k) in (None, ""):
                    body[k] = v
            if not body:
                continue
            body["imported_at"] = now
            op = {"op": "update"}
            if ex.get("version"):
                op["if_version"] = ex["version"]
            else:
                missing_versions.append(doc_id)   # NEXUS refuses to change a record without its version
                continue
            updated.append(o.get("title", doc_id))
        p = os.path.join(a.out, "docs", f"t_{len(writes):03d}.json")
        json.dump(body, open(p, "w", encoding="utf-8"), ensure_ascii=False)
        writes.append({**op, "collection": "tenders", "doc_id": doc_id, "file_path": os.path.abspath(p)})

    status = {"last_sync": now, "added": len(added), "updated": len(updated), "deadline_changes": deadline_changes[:30],
              "workbook_rows": len(wb_rows), "by": "NEXUS auto-sync"}
    sc = scanner_status(a.last_scan)
    if sc:
        status["scanner"] = sc
    p = os.path.join(a.out, "docs", "sync_status.json")
    json.dump(status, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    w = {"op": "set", "collection": "system", "doc_id": "sync", "file_path": os.path.abspath(p)}
    if a.status_version:
        w["if_version"] = a.status_version
    writes.append(w)

    if a.staff_dir:
        staff = []
        for f in glob.glob(os.path.join(a.staff_dir, "staff", "*.json")):
            d = json.load(open(f, encoding="utf-8")); d = d.get("data", d)
            staff.append({k: d.get(k, "") for k in ("name", "position", "email", "phone", "status", "n_email", "n_sms", "n_whatsapp", "alert_types")})
        json.dump({"updated": now, "staff": sorted(staff, key=lambda x: x["name"])}, open(os.path.join(a.out, "recipients.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for i in range(0, len(writes), 50):
        json.dump(writes[i:i + 50], open(os.path.join(a.out, f"writes_{i // 50:02d}.json"), "w", encoding="utf-8"), indent=0)
    print(json.dumps({"added": added, "updated": len(updated), "deadline_changes": deadline_changes,
                      "batches": (len(writes) + 49) // 50, "writes": len(writes),
                      "scanner_last_scan": sc and sc["last_scan"],
                      "skipped_no_version": missing_versions}, ensure_ascii=False))


if __name__ == "__main__":
    main()
